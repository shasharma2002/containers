#!/usr/bin/env python3
"""
Parse the commercial "LC OPEN AND CARGO IN TRANSIT" workbook into one flat
shipment record per sheet row, ready to join against carrier tracking data on
bl_number (see shipments.json in track_all.py).

    python read_sheet.py "LC OPEN AND CARGO IN TRANSIT.xlsx"
    python read_sheet.py path.xlsx --sheet "Sheet1" --json

SEVEN THINGS THIS SHEET WILL DO TO A NAIVE PARSER
--------------------------------------------------
1. LC/DAP sits immediately left of BL NO. and looks exactly as plausible as
   a BL number. We never guess: BL NO. is located by matching its header
   text explicitly, and we fail loudly (SystemExit) if that header can't be
   found at all -- silently reading the wrong column is worse than crashing.

2. SHIPPING LINE, not the BL prefix, tells us the carrier. Maersk's BLs on
   this sheet are bare 9-digit numbers with no prefix at all (270383678),
   so BIC-prefix classification (carrier_id.classify) cannot work on them
   regardless of how good the prefix table is. carrier_id.classify is a
   fallback for track_all.py to use ONLY when SHIPPING LINE is blank or
   unrecognised -- this module never calls it.

3. Bulk cargo rows (a tanker name like "MT WECO SENNA" in SHIPPING LINE, no
   BL at all) are not container shipments. They're skipped silently and
   counted, never reported as errors.

4. HAPAG has no adapter anywhere in this pipeline. Rows shipping on
   Hapag-Lloyd are real shipments, so they still come out of this module
   with shipping_line_raw="HAPAG" and carrier=None -- it's track_all.py's
   job to recognise that specific case and record it as unsupported_carrier
   rather than attempting a lookup or falling back to prefix guessing.

5. The same BL can appear on more than one sheet row (split shipments).
   This module does NOT deduplicate -- every sheet row is real commercial
   data and belongs in shipments.json as its own row. Deduplicating for the
   actual carrier lookups (one fetch per unique BL) is track_all.py's job,
   done once it has the full list.

6. Continuation rows -- an extra commodity/qty line under the same
   shipment -- have a blank BENEFICIARY *and* a blank BL NO. We forward-fill
   the shipment-identity fields (BL, shipping line, beneficiary, broker,
   PI no., commercial terms, ETA) from the block's first row, but leave
   commodity/qty/amount alone since those are genuinely per-line.
   A bulk-cargo row is NOT a continuation row: it has its own BENEFICIARY
   filled in, just no BL. That distinction is what tells the two apart.

7. The real header row isn't necessarily row 1 -- there can be merged
   title/blank rows above it. We scan for the row that actually contains a
   BL NO. header instead of assuming a fixed offset.
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

HEADER_SCAN_ROWS = 15

# Header text, normalised (letters/digits only, uppercased), that identifies
# the BL NO. column. This is the one header we refuse to guess about.
BLNO_ALIASES = {"BLNO", "BLNUMBER", "BILLOFLADINGNO", "BILLOFLADING", "BLNOS"}

# Every other header we know how to use. LC/DAP is captured under its own
# field (lc_dap) purely so it's obviously never mistaken for bl_number_raw --
# it is not part of the shipment record this module returns.
HEADER_ALIASES: dict[str, str] = {
    "BENEFICIARY": "beneficiary",
    "BROKER": "broker",
    "COMMODITIES": "commodity",
    "COMMODITY": "commodity",
    "PINO": "pi_no",
    "PINUMBER": "pi_no",
    "QTY": "qty",
    "QUANTITY": "qty",
    "AMOUNT": "amount",
    "DELIVERY": "delivery_terms",
    "FREEDAYS": "free_days",
    "PAYMENT": "payment",
    "LCDAP": "lc_dap",
    "SHIPPINGLINE": "shipping_line_raw",
    "ETA": "eta_raw",
}
for _alias in BLNO_ALIASES:
    HEADER_ALIASES[_alias] = "bl_number_raw"

# SHIPPING LINE -> our carrier code. Deliberately small and exact-match only;
# anything not in here (blank, a tanker name, "HAPAG") comes back as carrier
# None and is track_all.py's decision, not this module's.
CARRIER_MAP = {"MAERSK": "MAEU", "MSC": "MSCU"}

ETA_RE = re.compile(r"^\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})\s*(.*)$")


class SheetError(RuntimeError):
    pass


# --------------------------------------------------------------------------- #
# header location
# --------------------------------------------------------------------------- #

def _norm(h) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(h if h is not None else "").strip().upper())


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Scan the first HEADER_SCAN_ROWS rows for one containing a BL NO. header.

    Returns (header_row_index, {field_name: column_index}). Raises SheetError
    if no row in that window has a recognisable BL NO. header -- we do not
    fall back to a fixed row number or a fuzzy guess.
    """
    for r_idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        normed = {_norm(c): c_idx for c_idx, c in enumerate(row) if c is not None}
        if not any(alias in normed for alias in BLNO_ALIASES):
            continue
        col_map = {}
        for norm_h, c_idx in normed.items():
            field = HEADER_ALIASES.get(norm_h)
            if field:
                col_map[field] = c_idx
        return r_idx, col_map

    raise SheetError(
        f"No 'BL NO.' header found in the first {HEADER_SCAN_ROWS} rows. "
        "Refusing to guess which column is the BL -- pass --sheet if the "
        "data is on a different tab, or check the file by hand."
    )


# --------------------------------------------------------------------------- #
# cell cleanup
# --------------------------------------------------------------------------- #

def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _clean(v):
    """Strip/blank-out strings; leave numbers (qty, amount) alone."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    if isinstance(v, (datetime.datetime, datetime.date)):
        # Seen in the wild: a PI NO. cell holding a raw Excel date. Not
        # JSON-serialisable as-is, and a date is not useful as a PI number
        # anyway -- surface it as text rather than crash downstream.
        return v.isoformat()
    return v


def _bl_string(v) -> str | None:
    """Normalise a BL cell to a query string.

    Maersk BLs on this sheet are bare numbers (270383678) with no letter
    prefix, and Excel often stores a numeric-looking cell as an actual
    float/int -- str(270383678.0) is "270383678.0", which is wrong. Handle
    that before falling through to the generic string path.
    """
    if v is None:
        return None
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    s = str(v).strip().upper().replace(" ", "")
    return s or None


def _parse_eta(raw: str | None) -> tuple[str | None, str | None]:
    """'29-08-2026 DRYPORT' -> ('2026-08-29', 'DRYPORT')."""
    if not raw:
        return None, None
    raw = str(raw).strip()
    if not raw:
        return None, None
    m = ETA_RE.match(raw)
    if not m:
        return None, raw
    d, mo, y, place = m.groups()
    try:
        date = f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"
    except ValueError:
        return None, raw
    return date, (place.strip() or None)


def classify_shipping_line(shipping_line_raw: str | None) -> str | None:
    """SHIPPING LINE text -> our carrier code, or None.

    None covers three different situations on purpose -- blank, a genuinely
    unrecognised value, and a KNOWN-but-unsupported carrier like HAPAG.
    Telling those apart (fallback to carrier_id.classify vs. an
    unsupported_carrier error) is track_all.py's job, using the
    shipping_line_raw field this function leaves untouched.
    """
    if not shipping_line_raw:
        return None
    return CARRIER_MAP.get(shipping_line_raw.strip().upper())


# --------------------------------------------------------------------------- #
# parsing
# --------------------------------------------------------------------------- #

def get_shipments(path: Path, sheet: str | None = None) -> dict:
    """-> {"shipments": [...], "skipped_bulk": n, "skipped_no_bl": n, "skipped_blank": n}"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        raise SheetError("Empty sheet.")

    header_idx, col_map = _find_header(rows)

    shipments: list[dict] = []
    skipped_bulk = skipped_no_bl = skipped_blank = 0
    context: dict | None = None   # last block's shipment-identity fields

    for row_number, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        beneficiary = _clean(_cell(row, col_map.get("beneficiary")))
        bl_raw = _cell(row, col_map.get("bl_number_raw"))
        commodity = _clean(_cell(row, col_map.get("commodity")))
        qty = _clean(_cell(row, col_map.get("qty")))
        amount = _clean(_cell(row, col_map.get("amount")))
        bl_number = _bl_string(bl_raw)

        if not any([beneficiary, bl_number, commodity, qty, amount]):
            skipped_blank += 1
            continue

        if not bl_number:
            if beneficiary:
                # New block, no BL at all -- bulk/vessel cargo, not a
                # container shipment. Skip silently; this is not an error.
                skipped_bulk += 1
                continue
            if context is None:
                # Blank BL with nothing above to inherit from.
                skipped_no_bl += 1
                continue
            identity = context
        else:
            identity = {
                "bl_number": bl_number,
                "shipping_line_raw": _clean(_cell(row, col_map.get("shipping_line_raw"))),
                "beneficiary": beneficiary,
                "broker": _clean(_cell(row, col_map.get("broker"))),
                "pi_no": _clean(_cell(row, col_map.get("pi_no"))),
                "delivery_terms": _clean(_cell(row, col_map.get("delivery_terms"))),
                "free_days": _clean(_cell(row, col_map.get("free_days"))),
                "eta_raw": _clean(_cell(row, col_map.get("eta_raw"))),
            }
            context = identity

        carrier = classify_shipping_line(identity["shipping_line_raw"])
        eta_date, eta_place = _parse_eta(identity["eta_raw"])

        shipments.append({
            "row_number": row_number,
            "bl_number": identity["bl_number"],
            "carrier": carrier,
            "shipping_line_raw": identity["shipping_line_raw"],
            "beneficiary": identity["beneficiary"],
            "broker": identity["broker"],
            "commodity": commodity,
            "pi_no": identity["pi_no"],
            "qty": qty,
            "amount": amount,
            "delivery_terms": identity["delivery_terms"],
            "free_days": identity["free_days"],
            "sheet_eta_raw": identity["eta_raw"],
            "sheet_eta_date": eta_date,
            "sheet_eta_place": eta_place,
        })

    return {
        "shipments": shipments,
        "skipped_bulk": skipped_bulk,
        "skipped_no_bl": skipped_no_bl,
        "skipped_blank": skipped_blank,
    }


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main() -> int:
    ap = argparse.ArgumentParser(
        description="Parse the LC OPEN AND CARGO IN TRANSIT workbook.")
    ap.add_argument("path", type=Path)
    ap.add_argument("--sheet")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    try:
        result = get_shipments(args.path, args.sheet)
    except SheetError as exc:
        print(f"\nFAILED: {exc}", file=sys.stderr)
        return 1

    shipments = result["shipments"]

    if args.json:
        print(json.dumps(shipments, indent=2, ensure_ascii=False))
        return 0

    print(f"\n{len(shipments)} shipment row(s) parsed  "
          f"(skipped: {result['skipped_bulk']} bulk cargo, "
          f"{result['skipped_no_bl']} no BL / no block, "
          f"{result['skipped_blank']} blank)\n")
    for s in shipments:
        print(f"row {s['row_number']:<5} BL {s['bl_number']:<20} "
              f"carrier {s['carrier'] or '?':<6} ({s['shipping_line_raw'] or '-'})  "
              f"{(s['beneficiary'] or '')[:24]:<24} {(s['commodity'] or '')[:16]:<16} "
              f"qty={s['qty']}  ETA {s['sheet_eta_date'] or '-'} @ {s['sheet_eta_place'] or '-'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
